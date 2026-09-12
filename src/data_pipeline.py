from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from docx import Document

from .config import (
    BLEND_TOLERANCE_HIGH,
    BLEND_TOLERANCE_LOW,
    BTE_BASE_RANGE_MAX,
    BTE_BASE_RANGE_MIN,
    BTE_CALCULATION_CR_THRESHOLD,
    COMPOSITION_FEATURES,
    ENGINE_FEATURES,
    MODEL_FEATURES,
    PERFORMANCE_TARGETS,
    PROPERTY_TARGETS,
)


def _column_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


COLUMN_MAP = {
    "srno": "source_serial_no",
    "gasoline": "gasoline_pct",
    "ethanol": "ethanol_pct",
    "pentanol": "pentanol_pct",
    "propenol": "propanol_pct",
    "propanol": "propanol_pct",
    "butanol": "butanol_pct",
    "nmythylanniline": "n_methylaniline_pct",
    "nmethylaniline": "n_methylaniline_pct",
    "torquenm": "torque_nm",
    "speedrpm": "speed_rpm",
    "compressionratio": "compression_ratio",
    "cvmjkg": "cv_mj_kg",
    "af": "air_fuel_ratio",
    "viscositym2s": "viscosity_m2_s",
    "ron": "ron",
    "oxugen": "oxygen_pct",
    "oxygen": "oxygen_pct",
    "bte": "bte_pct",
    "bsfcgkwh": "bsfc_g_kwh",
    "covol": "co_vol_pct",
    "hcppm": "hc_ppm",
}


def _canonical_name(original: object) -> str:
    key = _column_key(original)
    for token, canonical in COLUMN_MAP.items():
        if key == token or key.startswith(token):
            return canonical
    return key


def _json_number(value: Any) -> Any:
    if pd.isna(value):
        return None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    return value


def load_and_clean_experimental_data(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    """Load the first worksheet, quarantine invalid records, and return ML-ready data.

    Composition records are accepted only when their total is 99.99-100.01%.
    Tiny accepted rounding deviations are rescaled to exactly 100%. Larger
    deviations are retained in the quarantine file, never silently deleted.
    """

    source_path = Path(path)
    raw = pd.read_excel(source_path, sheet_name=0)
    original_shape = raw.shape
    raw.columns = [_canonical_name(c) for c in raw.columns]
    raw["source_excel_row"] = np.arange(len(raw)) + 2

    required = MODEL_FEATURES + PROPERTY_TARGETS + PERFORMANCE_TARGETS
    missing_columns = sorted(set(required) - set(raw.columns))
    if missing_columns:
        raise ValueError(f"Missing required columns: {missing_columns}")

    source_nonempty = ~raw[required].isna().all(axis=1)
    working = raw.loc[source_nonempty].copy()
    for column in required + ["source_serial_no"]:
        if column in working:
            working[column] = pd.to_numeric(working[column], errors="coerce")

    missing_required = working[required].isna().any(axis=1)
    missing_quarantine = working.loc[missing_required].copy()
    missing_quarantine["quarantine_reason"] = "missing_or_non_numeric_required_value"
    working = working.loc[~missing_required].copy()

    working["composition_sum_original_pct"] = working[COMPOSITION_FEATURES].sum(axis=1)
    valid_sum = working["composition_sum_original_pct"].between(
        BLEND_TOLERANCE_LOW - 1e-9,
        BLEND_TOLERANCE_HIGH + 1e-9,
    )
    blend_quarantine = working.loc[~valid_sum].copy()
    blend_quarantine["quarantine_reason"] = "fuel_composition_not_within_99.99_to_100.01_pct"
    working = working.loc[valid_sum].copy()

    scale = 100.0 / working["composition_sum_original_pct"]
    working.loc[:, COMPOSITION_FEATURES] = working[COMPOSITION_FEATURES].mul(scale, axis=0)
    working["composition_sum_pct"] = working[COMPOSITION_FEATURES].sum(axis=1)
    # Group nominally identical blends together even when the source total is
    # 99.99 because of reporting precision. This is deliberately coarser than
    # the stored normalized values and prevents near-duplicate blend leakage.
    working["blend_group"] = working[COMPOSITION_FEATURES].round(1).astype(str).agg("|".join, axis=1)

    duplicate_mask = working[required].duplicated(keep="first")
    duplicate_quarantine = working.loc[duplicate_mask].copy()
    duplicate_quarantine["quarantine_reason"] = "exact_duplicate_model_record"
    working = working.loc[~duplicate_mask].copy()

    quarantine = pd.concat(
        [missing_quarantine, blend_quarantine, duplicate_quarantine],
        ignore_index=True,
        sort=False,
    )
    if not quarantine.empty:
        quarantine = quarantine.sort_values("source_excel_row").reset_index(drop=True)

    all_empty_rows = int((~source_nonempty).sum())
    numeric_compositions = raw.loc[source_nonempty, COMPOSITION_FEATURES].apply(
        pd.to_numeric, errors="coerce"
    )
    composition_counts = (
        numeric_compositions.sum(axis=1, min_count=len(COMPOSITION_FEATURES))
        .round(2)
        .value_counts(dropna=False)
        .sort_index()
    )
    high_cr_bte = working.loc[
        working["compression_ratio"] >= BTE_CALCULATION_CR_THRESHOLD,
        "bte_pct",
    ]
    high_cr_envelopes = {
        str(float(compression_ratio)): {
            "rows": int(len(group)),
            "source_min_pct": float(group["bte_pct"].min()),
            "source_max_pct": float(group["bte_pct"].max()),
        }
        for compression_ratio, group in working.loc[
            working["compression_ratio"] >= BTE_CALCULATION_CR_THRESHOLD
        ].groupby("compression_ratio", sort=True)
    }
    audit = {
        "source_file": source_path.name,
        "source_rows": int(original_shape[0]),
        "source_columns": int(original_shape[1]),
        "completely_blank_rows": all_empty_rows,
        "nonempty_source_rows": int(source_nonempty.sum()),
        "quarantined_rows": int(len(quarantine)),
        "valid_model_rows": int(len(working)),
        "valid_blend_groups": int(working["blend_group"].nunique()),
        "blend_group_rounding_pct": 0.1,
        "exact_duplicate_model_rows_removed": int(len(duplicate_quarantine)),
        "composition_sum_counts": {str(_json_number(k)): int(v) for k, v in composition_counts.items()},
        "training_ranges": {
            column: {
                "min": float(working[column].min()),
                "max": float(working[column].max()),
            }
            for column in MODEL_FEATURES + PROPERTY_TARGETS + PERFORMANCE_TARGETS
        },
        "bte_cr_ge_9_5": {
            "rows": int(len(high_cr_bte)),
            "source_min_pct": float(high_cr_bte.min()),
            "source_max_pct": float(high_cr_bte.max()),
            "source_envelopes_by_cr": high_cr_envelopes,
            "base_output_min_pct": BTE_BASE_RANGE_MIN,
            "base_output_max_pct": BTE_BASE_RANGE_MAX,
            "normalization_formula": "q = clip((raw_bte - source_min_at_cr) / (source_max_at_cr - source_min_at_cr), 0, 1)",
            "output_formula": "final_bte = blend_range_min + q * (blend_range_max - blend_range_min)",
            "outside_26_to_36_rows": int(
                (
                    (working["compression_ratio"] >= BTE_CALCULATION_CR_THRESHOLD)
                    & ~working["bte_pct"].between(
                        BTE_BASE_RANGE_MIN,
                        BTE_BASE_RANGE_MAX,
                    )
                ).sum()
            ),
        },
        "notes": [
            "Quarantined records remain available for manual engineering review.",
            "BTE source values are preserved. For CR >= 9.5, the internal deep-network BTE signal is normalized within the CR-specific cleaned-data envelope and mapped into the active user-defined blend range.",
            "User-defined blend ranges are kept separate from raw model validation because they are engineering priors rather than measured target labels.",
            "Fuel properties are estimated from composition for new blends and are not accepted as arbitrary user inputs.",
            "Blend groups use 0.1 percentage-point rounding to keep 99.99% reporting noise inside one validation group.",
        ],
    }

    # Retain source order while putting model fields first.
    ordered = (
        ["source_excel_row", "source_serial_no"]
        + MODEL_FEATURES
        + PROPERTY_TARGETS
        + PERFORMANCE_TARGETS
        + ["composition_sum_original_pct", "composition_sum_pct", "blend_group"]
    )
    working = working[ordered].sort_values("source_excel_row").reset_index(drop=True)
    return working, quarantine, audit


def _extract_ethanol_default(blend_text: str) -> float:
    values = [float(x) for x in re.findall(r"\bE(\d+(?:\.\d+)?)\b", blend_text.upper())]
    return max(values) if values else 0.0


def extract_vehicle_database(path: str | Path) -> pd.DataFrame:
    """Extract the supplied vehicle comparison table without inventing specifications."""

    document = Document(Path(path))
    if not document.tables:
        raise ValueError("No vehicle table was found in the supplied DOCX.")

    table = document.tables[0]
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    headers = rows[0]
    records = []
    for values in rows[1:]:
        record = dict(zip(headers, values))
        if not any(str(v).strip() for v in record.values()):
            continue
        records.append(record)

    vehicles = pd.DataFrame(records)
    rename = {
        "Company": "company",
        "Model": "model",
        "Compression Ratio": "compression_ratio",
        "Category": "category",
        "Recommended Fuel (RON)": "recommended_fuel_ron",
        "Fuel Blend (Octane)": "fuel_blend_octane",
        "Year of Launch": "year_of_launch",
        "BS Norm (at Launch)": "bs_norm_at_launch",
        "Estimated Mileage on Blend (kmpl)": "estimated_mileage_on_blend_kmpl",
    }
    vehicles = vehicles.rename(columns=rename)
    expected = list(rename.values())
    if not set(expected).issubset(vehicles.columns):
        raise ValueError("The vehicle table structure did not match the expected supplied specification.")

    vehicles["compression_ratio"] = pd.to_numeric(vehicles["compression_ratio"], errors="coerce")
    vehicles["year_of_launch"] = pd.to_numeric(vehicles["year_of_launch"], errors="coerce").astype("Int64")
    vehicles["default_ethanol_pct"] = vehicles["fuel_blend_octane"].map(_extract_ethanol_default)
    vehicles["default_gasoline_pct"] = 100.0 - vehicles["default_ethanol_pct"]

    # Torque and RPM are absent from the supplied vehicle table. These are app
    # starting points based on the experimental dataset, not vehicle claims.
    vehicles["default_torque_nm"] = 5.51
    vehicles["default_speed_rpm"] = 2500.0
    vehicles["operating_default_note"] = "Dataset starting point - verify the intended engine test condition"
    vehicles["vehicle_id"] = vehicles["company"].str.strip() + " | " + vehicles["model"].str.strip()

    columns = [
        "vehicle_id",
        "company",
        "model",
        "compression_ratio",
        "category",
        "recommended_fuel_ron",
        "fuel_blend_octane",
        "year_of_launch",
        "bs_norm_at_launch",
        "estimated_mileage_on_blend_kmpl",
        "default_gasoline_pct",
        "default_ethanol_pct",
        "default_torque_nm",
        "default_speed_rpm",
        "operating_default_note",
    ]
    return vehicles[columns].reset_index(drop=True)


def write_audit(audit: dict[str, Any], path: str | Path) -> None:
    Path(path).write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")
